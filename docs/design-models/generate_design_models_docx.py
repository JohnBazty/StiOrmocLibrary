"""Generate vertical Mermaid design-model figures into a Word document."""
from __future__ import annotations

import base64
import io
import urllib.error
import urllib.request
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT_DIR = Path(r"c:\Users\Ethan\Documents\github\StiOrmocLibrary\docs\design-models")
DOWNLOADS = Path(r"c:\Users\Ethan\Downloads")
OUT_DIR.mkdir(parents=True, exist_ok=True)

FIGURES: list[tuple[str, str, str]] = [
    (
        "Figure 1. Overall System Design Model (Users and Admin)",
        "overall_system",
        """flowchart TB
  Student[Student]
  Faculty[Faculty]
  Admin[Admin_Librarian]
  UserPortal[Student_Faculty_Web_Portal]
  AdminPortal[Admin_Web_Portal]
  API[Node_Express_API]
  DB[(Supabase_Postgres)]
  Storage[Supabase_Storage]
  Student --> UserPortal
  Faculty --> UserPortal
  Admin --> AdminPortal
  UserPortal -->|JWT_REST| API
  AdminPortal -->|JWT_REST| API
  API --> DB
  API --> Storage""",
    ),
    (
        "Figure 2. User Features Design Model (Student and Faculty)",
        "user_features",
        """flowchart TB
  ULogin[Login]
  UReg[Student_Register]
  UPortal[User_Portal]
  UDash[Dashboard_Overview]
  UCat[Book_Catalog]
  UFloor[Library_Floor_Plan]
  UCart[Book_Cart]
  UResearch[Research_Thesis_Catalog]
  UHist[Borrowing_History]
  URes[My_Reservations]
  UPrint[Printing_Service_Student]
  UQR[QR_Attendance_Pass]
  UNotif[Notifications]
  UFine[Fines_and_Receipts]
  UClear[Clearance_Status]
  ULogin --> UPortal
  UReg --> ULogin
  UPortal --> UDash
  UDash --> UCat
  UDash --> UFloor
  UDash --> UCart
  UDash --> UResearch
  UDash --> UHist
  UDash --> URes
  UDash --> UPrint
  UDash --> UQR
  UDash --> UNotif
  UDash --> UFine
  UDash --> UClear
  UCat --> UCart
  UCat --> UFloor
  UCart --> UHist
  UCat --> URes
  UResearch --> URes
  UHist --> UFine
  UHist --> UClear
  UPrint --> UFine
  UFine --> UClear""",
    ),
    (
        "Figure 3. User Borrowing Process Design Model",
        "user_borrowing",
        """flowchart TB
  B1[Browse_Book_Catalog]
  B2[View_Book_Details]
  B3[Add_to_Book_Cart]
  B4[Submit_Borrow_Request]
  B5[Pending_Request]
  B6[Admin_Issues_Book]
  B7[Borrowed_Active]
  B8[Due_Date_Tracking]
  B9[Return_at_Circulation_Desk]
  B10[Returned]
  B11[Overdue_Fine_if_Late]
  B12[Cancel_Pending_Request]
  B1 --> B2
  B2 --> B3
  B3 --> B4
  B4 --> B5
  B5 --> B6
  B5 --> B12
  B6 --> B7
  B7 --> B8
  B8 --> B9
  B9 --> B10
  B8 --> B11
  B11 --> B10""",
    ),
    (
        "Figure 4. User Reservation Process Design Model",
        "user_reservation",
        """flowchart TB
  R1[Title_Unavailable]
  R2[Join_Reservation_Queue]
  R3[Queued]
  R4[Admin_Approves_Ready]
  R5[Ready_for_Pickup]
  R6[User_Claims_at_Desk]
  R7[Borrowed]
  R8[Cancel_Reservation]
  R1 --> R2
  R2 --> R3
  R3 --> R4
  R3 --> R8
  R4 --> R5
  R5 --> R6
  R6 --> R7
  R5 --> R8""",
    ),
    (
        "Figure 5. User QR Attendance Process Design Model",
        "user_attendance",
        """flowchart TB
  A1[Open_QR_Attendance]
  A2[View_Permanent_Pass]
  A3[Download_QR_Image]
  A4[Present_Pass_Offline]
  A5[Staff_Scans_QR]
  A6[Check_In_or_Check_Out]
  A7[Attendance_History_Updated]
  A1 --> A2
  A2 --> A3
  A2 --> A4
  A3 --> A4
  A4 --> A5
  A5 --> A6
  A6 --> A7""",
    ),
    (
        "Figure 6. User Printing Process Design Model (Student)",
        "user_printing",
        """flowchart TB
  P1[Open_Printing_Service]
  P2[Upload_PDF_or_DOCX]
  P3[Auto_Page_Count_and_Quote]
  P4[Submit_Print_Request]
  P5[Pending_Unpaid]
  P6[Pay_Cash_at_Counter]
  P7[Admin_Records_Payment]
  P8[Printing]
  P9[Ready_for_Pickup]
  P10[Completed]
  P11[Download_Digital_Receipt]
  P12[Cancel_if_Still_Unpaid]
  P1 --> P2
  P2 --> P3
  P3 --> P4
  P4 --> P5
  P5 --> P6
  P5 --> P12
  P6 --> P7
  P7 --> P8
  P8 --> P9
  P9 --> P10
  P7 --> P11""",
    ),
    (
        "Figure 7. User Account Support Design Model",
        "user_account",
        """flowchart TB
  S1[User_Dashboard]
  S2[Borrowing_History]
  S3[Notifications]
  S4[Fines_Module]
  S5[Clearance_Status]
  S6[View_Balances_and_Receipts]
  S7[Cleared_or_Blocked]
  S1 --> S2
  S1 --> S3
  S1 --> S4
  S1 --> S5
  S2 --> S4
  S4 --> S6
  S4 --> S5
  S5 --> S7
  S3 --> S3""",
    ),
    (
        "Figure 8. Admin Features Design Model",
        "admin_features",
        """flowchart TB
  AdminLogin[Admin_Login]
  AdminPortal[Admin_Portal]
  ADash[Dashboard]

  Ops[Operations]
  Res[Resources]
  People[People_and_Records]

  ACat[Books_and_Research]
  AArch[Book_Archive]
  ACatg[Categories]
  ACirc[Borrow_and_Return]
  AResQ[Reservations_Queue]
  AFines[Fines]

  AInv[Inventory]
  AFloor[Floor_Plan]
  APrintQ[Printing_Queue]
  ASupp[Print_Supplies]

  AAtt[Attendance]
  AUsers[Users_MIS]
  AClear[Clearance]
  AAnn[Announcements]
  ARep[Reports]

  AdminLogin --> AdminPortal
  AdminPortal --> ADash
  ADash --> Ops
  ADash --> Res
  ADash --> People

  Ops --> ACat
  ACat --> AArch
  AArch --> ACatg
  ACatg --> ACirc
  ACirc --> AResQ
  AResQ --> AFines

  Res --> AInv
  AInv --> AFloor
  AFloor --> APrintQ
  APrintQ --> ASupp

  People --> AAtt
  AAtt --> AUsers
  AUsers --> AClear
  AClear --> AAnn
  AAnn --> ARep""",
    ),
    (
        "Figure 9. Admin Circulation Design Model",
        "admin_circulation",
        """flowchart TB
  C1[Open_Borrow_and_Return]
  C2[Scan_User_or_Barcode]
  C3{Transaction_Type}
  C4[Issue_Checkout]
  C5[Mark_Borrowed]
  C6[Process_Return]
  C7{Overdue}
  C8[Create_or_Update_Fine]
  C9[Close_Loan_Returned]
  C10[Monitor_Active_Loans]
  C1 --> C2
  C2 --> C3
  C3 -->|Issue| C4
  C4 --> C5
  C3 -->|Return| C6
  C6 --> C7
  C7 -->|Yes| C8
  C7 -->|No| C9
  C8 --> C9
  C5 --> C10
  C9 --> C10""",
    ),
    (
        "Figure 10. Admin Catalog and Inventory Design Model",
        "admin_catalog_inventory",
        """flowchart TB
  I1[Books_and_Research_Module]
  I2[Add_or_Edit_Title]
  I3[Register_Physical_Copies]
  I4[Barcode_ISBN_Scan]
  I5[Manage_Categories]
  I6[Inventory_Dashboard]
  I7[Update_Condition_Availability]
  I8[Thesis_Inventory]
  I9[Export_CSV_PDF]
  I10[Book_Archive]
  I1 --> I2
  I1 --> I3
  I1 --> I4
  I1 --> I5
  I1 --> I6
  I6 --> I7
  I6 --> I8
  I6 --> I9
  I1 --> I10""",
    ),
    (
        "Figure 11. Admin Printing and Supplies Design Model",
        "admin_printing",
        """flowchart TB
  Pr1[Printing_Queue]
  Pr2[Pause_or_Resume_Service]
  Pr3[Select_Print_Request]
  Pr4[Record_Cash_Payment]
  Pr5[Update_Job_Status]
  Pr6[Printing]
  Pr7[Ready_for_Pickup]
  Pr8[Completed]
  Pr9[Print_Supplies]
  Pr10[Update_Ink_Stock]
  Pr11[Update_Paper_Stock]
  Pr12[Review_Revenue]
  Pr1 --> Pr2
  Pr1 --> Pr3
  Pr3 --> Pr4
  Pr4 --> Pr5
  Pr5 --> Pr6
  Pr6 --> Pr7
  Pr7 --> Pr8
  Pr4 --> Pr9
  Pr9 --> Pr10
  Pr9 --> Pr11
  Pr9 --> Pr12""",
    ),
    (
        "Figure 12. Admin People and Records Design Model",
        "admin_people",
        """flowchart TB
  P1[Users_MIS]
  P2[Create_or_Edit_Account]
  P3[Assign_Role]
  P4[Activate_Deactivate_Archive]
  P5[Attendance_Monitoring]
  P6[Scan_Student_Faculty_QR]
  P7[Update_Visit_Logs]
  P8[Clearance_Module]
  P9[Review_Standing]
  P10[Authorized_Override]
  P11[Announcements]
  P12[Publish_Notice]
  P13[Reports]
  P1 --> P2
  P2 --> P3
  P3 --> P4
  P5 --> P6
  P6 --> P7
  P8 --> P9
  P9 --> P10
  P11 --> P12
  P13 --> P13""",
    ),
    (
        "Figure 13. User and Admin Feature Interaction Design Model",
        "user_admin_interaction",
        """flowchart TB
  UCat[User_Catalog_Cart_Research]
  UBorrow[User_Borrow_Reserve_History]
  UPrint[User_Print_QR_Notifications]
  UAcct[User_Fines_Clearance]
  ACat[Admin_Catalog_Categories_Archive]
  ACirc[Admin_Circulation_Reservations]
  AOps[Admin_Printing_Supplies_Attendance]
  ARec[Admin_Fines_Clearance_Users_Reports]
  UCat --> ACat
  ACat --> UCat
  UBorrow --> ACirc
  ACirc --> UBorrow
  UPrint --> AOps
  AOps --> UPrint
  UAcct --> ARec
  ARec --> UAcct""",
    ),
]


def mermaid_to_url(source: str) -> str:
    encoded = base64.urlsafe_b64encode(source.encode("utf-8")).decode("ascii")
    return f"https://mermaid.ink/img/{encoded}?type=png"


def fetch_png(source: str) -> bytes | None:
    url = mermaid_to_url(source)
    req = urllib.request.Request(url, headers={"User-Agent": "SmartLib-DesignModels/1.0"})
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            data = resp.read()
            if data[:8] == b"\x89PNG\r\n\x1a\n":
                return data
    except Exception as exc:  # noqa: BLE001
        print(f"Render failed: {exc}")
    return None


def set_run_font(run, name="Times New Roman", size=12, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_heading_styled(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=14 if level == 1 else 12, bold=True)
    return p


def build_docx(images: dict[str, Path]) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("STI Ormoc Smart Library")
    set_run_font(run, size=16, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("System Design Models\nUser and Admin Feature Flows")
    set_run_font(run, size=14, bold=True)

    note = doc.add_paragraph()
    run = note.add_run(
        "All diagrams are drawn top-to-bottom (vertical) for manuscript layout. "
        "They reflect the implemented React web portals and Node.js API modules."
    )
    set_run_font(run, size=11)
    run.italic = True

    add_heading_styled(doc, "1. Design Model Figures", level=1)

    for caption, key, source in FIGURES:
        add_heading_styled(doc, caption, level=2)
        img = images.get(key)
        if img and img.exists():
            doc.add_picture(str(img), width=Inches(5.8))
            last = doc.paragraphs[-1]
            last.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p = doc.add_paragraph()
            run = p.add_run("[Diagram image could not be rendered automatically. Mermaid source is included below.]")
            set_run_font(run, size=10)
            run.font.color.rgb = RGBColor(0x99, 0x00, 0x00)

        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        set_run_font(run, size=11, bold=True)

        src_title = doc.add_paragraph()
        run = src_title.add_run("Mermaid source (vertical / flowchart TB):")
        set_run_font(run, size=10, bold=True)

        code = doc.add_paragraph()
        run = code.add_run(source.strip())
        set_run_font(run, name="Consolas", size=8)
        doc.add_paragraph()

    add_heading_styled(doc, "2. Scope Notes", level=1)
    bullets = [
        "User portal covers Student and Faculty. Printing Service is Student-only in the current build.",
        "Admin portal covers Librarian/System Administrator operational modules.",
        "Returning is performed at the admin circulation desk; users track status in Borrowing History.",
        "Diagrams align with Flowcharts.docx proposed processes and the implemented routes in apps/web.",
    ]
    for item in bullets:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=11)

    out_repo = OUT_DIR / "SmartLib_User_Admin_Design_Models.docx"
    out_dl = DOWNLOADS / "SmartLib_User_Admin_Design_Models.docx"
    doc.save(out_repo)
    doc.save(out_dl)
    return out_dl


def main():
    images: dict[str, Path] = {}
    for caption, key, source in FIGURES:
        mmd_path = OUT_DIR / f"{key}.mmd"
        mmd_path.write_text(source.strip() + "\n", encoding="utf-8")
        print(f"Rendering {key}...")
        png = fetch_png(source)
        if png:
            img_path = OUT_DIR / f"{key}.png"
            img_path.write_bytes(png)
            images[key] = img_path
            print(f"  OK -> {img_path.name} ({len(png)} bytes)")
        else:
            print(f"  WARN: no PNG for {key}")

    path = build_docx(images)
    print(f"DOCX written: {path}")
    print(f"Also saved: {OUT_DIR / 'SmartLib_User_Admin_Design_Models.docx'}")


if __name__ == "__main__":
    main()
